"""Lightweight parser adapters for the resource-constrained pilot server."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
import pypdf
from docx import Document as DocxDocument
from llama_index.core import Document


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".csv", ".md", ".txt", ".html", ".htm"}


class UnsupportedDocumentError(ValueError):
    pass


def parse_document(path: Path, original_filename: str) -> list[Document]:
    suffix = Path(original_filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocumentError(f"unsupported document format: {suffix or '<none>'}")

    if suffix == ".pdf":
        return _parse_pdf(path, original_filename)
    if suffix == ".docx":
        return _parse_docx(path, original_filename)
    if suffix in {".xlsx", ".csv"}:
        return _parse_tabular(path, original_filename, suffix)

    text = path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        text = re.sub(r"<script\b[^>]*>.*?</script>", " ", text, flags=re.I | re.S)
        text = re.sub(r"<style\b[^>]*>.*?</style>", " ", text, flags=re.I | re.S)
        text = re.sub(r"<[^>]+>", " ", text)
    return [Document(text=text, metadata={"filename": original_filename})]


def _parse_pdf(path: Path, original_filename: str) -> list[Document]:
    reader = pypdf.PdfReader(str(path))
    documents: list[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            documents.append(
                Document(
                    text=text,
                    metadata={"filename": original_filename, "page": page_number},
                )
            )
    if not documents:
        raise ValueError("PDF contains no extractable text; OCR parser is required")
    return documents


def _parse_docx(path: Path, original_filename: str) -> list[Document]:
    document = DocxDocument(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Document(text="\n".join(paragraphs), metadata={"filename": original_filename})]


def _parse_tabular(path: Path, original_filename: str, suffix: str) -> list[Document]:
    if suffix == ".csv":
        frame = pd.read_csv(path)
        sheets = {"csv": frame}
    else:
        sheets = pd.read_excel(path, sheet_name=None)
    documents: list[Document] = []
    for sheet_name, frame in sheets.items():
        text = frame.fillna("").astype(str).to_csv(index=False)
        documents.append(
            Document(
                text=text,
                metadata={"filename": original_filename, "sheet": str(sheet_name)},
            )
        )
    return documents
